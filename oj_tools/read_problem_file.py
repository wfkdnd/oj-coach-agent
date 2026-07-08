"""
读取 OJ 题目文本。

支持两种来源：
- 直接粘贴题目文本。
- 从 .txt / .md / .docx 文件读取题目。
"""

from __future__ import annotations

from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET


SUPPORTED_PROBLEM_EXTENSIONS = {".txt", ".md", ".docx"}
TEXT_ENCODINGS = ("utf-8", "utf-8-sig", "gb18030")
WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def read_problem(problem_text: str = "", file_path: str = "") -> str:
    """从粘贴文本或受支持的本地文件读取题目。

    problem_text 和 file_path 二选一即可。粘贴文本只会作为返回值传出，不会写入项目文件。
    """
    has_text = bool(problem_text.strip())
    has_file = bool(file_path.strip())

    if has_text and has_file:
        return "错误：problem_text 和 file_path 只能提供其中一个。"
    if has_text:
        return _normalize_text(problem_text)
    if has_file:
        return read_problem_file(file_path)
    return "错误：请提供 problem_text 或 file_path。"


def read_problem_file(file_path: str) -> str:
    """从 .txt、.md 或 .docx 文件读取 OJ 题目文本。"""
    path = Path(file_path).expanduser()
    error = _validate_problem_path(path)
    if error:
        return error

    if path.suffix.lower() == ".docx":
        return _read_docx(path)
    return _read_text_file(path)


def _validate_problem_path(path: Path) -> str:
    if not path.exists():
        return f"错误：文件不存在：{path}"
    if not path.is_file():
        return f"错误：路径不是文件：{path}"
    if path.suffix.lower() not in SUPPORTED_PROBLEM_EXTENSIONS:
        allowed = ", ".join(sorted(SUPPORTED_PROBLEM_EXTENSIONS))
        return f"错误：不支持的题目文件类型 '{path.suffix}'。支持：{allowed}"
    return ""


def _read_text_file(path: Path) -> str:
    last_error: Exception | None = None
    for encoding in TEXT_ENCODINGS:
        try:
            return _normalize_text(path.read_text(encoding=encoding))
        except UnicodeDecodeError as exc:
            last_error = exc
    return f"错误：文本文件解码失败：{path}。原因：{last_error}"


def _read_docx(path: Path) -> str:
    try:
        return _read_docx_with_python_docx(path)
    except ImportError:
        return _read_docx_with_stdlib(path)
    except Exception as exc:
        return f"错误：读取 docx 文档失败：{path}。原因：{exc}"


def _read_docx_with_python_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    chunks: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            chunks.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [_normalize_inline_text(cell.text) for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                chunks.append(line)

    return _normalize_text("\n".join(chunks))


def _read_docx_with_stdlib(path: Path) -> str:
    try:
        with zipfile.ZipFile(path) as archive:
            xml_bytes = archive.read("word/document.xml")
    except Exception as exc:
        return f"错误：读取 docx 文档内容失败：{path}。原因：{exc}"

    try:
        root = ET.fromstring(xml_bytes)
    except ET.ParseError as exc:
        return f"错误：解析 docx 文档 XML 失败：{path}。原因：{exc}"

    paragraph_tag = f"{{{WORD_NAMESPACE}}}p"
    text_tag = f"{{{WORD_NAMESPACE}}}t"
    lines: list[str] = []

    for paragraph in root.iter(paragraph_tag):
        text = "".join(node.text or "" for node in paragraph.iter(text_tag)).strip()
        if text:
            lines.append(text)

    return _normalize_text("\n".join(lines))


def _normalize_text(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n").strip()


def _normalize_inline_text(text: str) -> str:
    return " ".join(text.replace("\r\n", "\n").replace("\r", "\n").split())
